"""Avaliação da Band: passa as falas de referência do debate da Band pela sessão de checagem, com o
Claude de verdade, contra duas versões da base, e mede quando a publicação automática é segura.

  base congelada  só checagens com data de revisão anterior a 18/10/2026 21h59 (o que existia
                  antes do debate); registros sem data de revisão ficam de fora e são contados
  base completa   a base de checagens como está

Uso:
  python -P scripts/avaliacao_band.py rodar [--dados DADOS] [--falas F] [--saida PASTA] [--duble]
                                            [--aceitar-publicacao]
        congela a base, roda a sessão contra as duas versões e grava a planilha de marcação;
        --duble ensaia sem chamar o Claude; --aceitar-publicacao deixa entrar na base congelada
        também checagens só com data de publicação (UOL e BOL nunca trazem data de revisão)
  python -P scripts/avaliacao_band.py relatorio [--dados DADOS] [--saida PASTA] [--minimo 15]
        depois que o operador marcou a planilha: precisão por faixa de nota e o limiar escolhido

DADOS é --dados, a variável CHECAGEM_DADOS ou a pasta data/ do checkout principal. A base vem de
DADOS/checagens/; tudo o que a avaliação gera fica em DADOS/avaliacao/ (ou --saida):
  base_congelada/          base.json, vetores e congelamento.json (o que ficou de fora)
  rodada_congelada/        eventos.jsonl, fila.json, painel.json e agente/ (cada chamada ao Claude)
  rodada_completa/         idem, contra a base completa
  marcacao.xlsx            uma linha por ligação (fala × checagem); o operador marca certa/errada
  relatorio.md             precisão por faixa de nota em cada base e o limiar escolhido
  limiar.json              {"limiar": nota ou null, "origem": texto, "em": ISO}; a sessão da noite
                           do debate lê daqui (ausente ou null: publicação automática desligada)
"""
import json
import os
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from comum import BRASILIA, pasta_dados, raiz_principal  # noqa: E402
from sessao import le_falas, repassa  # noqa: E402

CORTE = datetime(2026, 10, 18, 21, 59, tzinfo=BRASILIA)  # início do debate da Band


def _ultimo_dia_inteiro_antes(corte):
    """As datas da base são dias UTC (como a fonte entrega) sem hora: só vale um dia que terminou
    inteiro antes do corte. O dia em que cai o corte é ambíguo e fica de fora."""
    return (corte.astimezone(timezone.utc).date() - timedelta(days=1)).isoformat()


def congela(origem, destino, corte=CORTE, origens=("revisao",)):
    """Grava em `destino` a base congelada derivada da base em `origem` (base.json e vetores):
    só checagens com data de revisão de um dia inteiro anterior ao corte. Com `origens` incluindo
    "publicacao", a data de publicação (a da URL, no caso do UOL) também vale. Devolve o resumo,
    que também fica em destino/congelamento.json."""
    from collections import Counter

    import numpy as np

    from base_checagens import le_base

    origem, destino = Path(origem), Path(destino)
    registros = le_base(origem)
    meta = json.loads((origem / "vetores.json").read_text(encoding="utf-8"))
    if meta["ids"] != [r["id"] for r in registros]:
        sys.exit(f"vetores fora de ordem com a base em {origem}: rode `base_checagens.py reconstruir`")
    X = np.load(origem / "vetores.npy")
    limite = _ultimo_dia_inteiro_antes(corte)
    com_data = [k for k, r in enumerate(registros) if r.get("data_origem") in origens and r.get("data")]
    ficam = [k for k in com_data if registros[k]["data"] <= limite]
    sem_data = set(range(len(registros))) - set(com_data)

    destino.mkdir(parents=True, exist_ok=True)
    congelados = [registros[k] for k in ficam]
    (destino / "base.json").write_text(json.dumps({"congelada_em": corte.isoformat(), "registros": congelados},
                                                  ensure_ascii=False, indent=1), encoding="utf-8")
    np.save(destino / "vetores.npy", X[ficam] if ficam else np.zeros((0, X.shape[1]), np.float32))
    (destino / "vetores.json").write_text(json.dumps({**meta, "ids": [r["id"] for r in congelados],
                                                      "frases": [meta["frases"][k] for k in ficam]},
                                                     ensure_ascii=False), encoding="utf-8")
    resumo = {
        "corte": corte.isoformat(),
        "regra": f"data de {' ou de '.join(origens)} (dia UTC) até {limite}; o dia do corte, sem hora, fica de fora",
        "base_completa": str(origem),
        "checagens_na_base_completa": len(registros),
        "checagens_na_base_congelada": len(ficam),
        "excluidas_sem_data_de_revisao": len(sem_data),
        "excluidas_sem_data_de_revisao_por_agencia": dict(Counter(
            registros[k].get("agencia") for k in sorted(sem_data)).most_common()),
        "excluidas_revisadas_depois_do_corte": len(com_data) - len(ficam),
    }
    (destino / "congelamento.json").write_text(json.dumps(resumo, ensure_ascii=False, indent=1), encoding="utf-8")
    return resumo


# ---------- rodadas ----------

BASES = ("congelada", "completa")


def rodada(falas, pasta_base, saida, chamar=None, paralelo=3, nome_base="", **kw):
    """Passa as falas [(quem, início em s, texto)] pela sessão contra a base em `pasta_base` e grava
    em `saida` o que sessao.repassa grava (eventos.jsonl, fila.json, painel.json e agente/). Uma rodada
    já terminada (com fila.json) não é refeita. kw: os de sessao.repassa (esperas_entre_tentativas)."""
    saida = Path(saida)
    if (saida / "fila.json").exists():
        print(f"{saida.name}: já rodada, aproveitando {saida / 'fila.json'}")
        return json.loads((saida / "fila.json").read_text(encoding="utf-8"))["resumo"]
    if (saida / "eventos.jsonl").exists():
        sys.exit(f"{saida} tem uma rodada interrompida (eventos.jsonl sem fila.json): apague a pasta e rode de novo")
    resumo = repassa(falas, pasta_base, saida, chamar=chamar, paralelo=paralelo, resumo={"base": nome_base}, **kw, meta={
        "debate": f"Debate do 2º turno · Band · 18/10/2026 · avaliação com a base {nome_base}",
        "encerrado": True,
    })
    print(json.dumps(resumo, ensure_ascii=False))
    return resumo


# ---------- planilha de marcação ----------

COLUNAS = ["minutagem", "quem", "fala", "trecho", "reformulação", "tipo", "agência", "data", "rótulo",
           "frase checada", "link", "nota congelada", "nota completa", "motivo", "marcação", "observação", "chave"]
MARCAS = ("certa", "errada")


def _chave(minutagem_s, url):
    return f"{minutagem_s:.1f} {url}"


def ligacoes(avaliacao):
    """As ligações (fala × checagem) propostas nas duas rodadas, uma vez cada, na ordem da fala:
    {chave: {campos..., "nota": {base: nota}, "motivo": {base: motivo}}}."""
    por_chave = {}
    for base in BASES:
        fila = json.loads((Path(avaliacao) / f"rodada_{base}" / "fila.json").read_text(encoding="utf-8"))["fila"]
        for lig in fila:
            for c in lig["checagens"]:
                k = _chave(lig["minutagem_s"], c["url"])
                linha = por_chave.setdefault(k, {
                    "minutagem_s": lig["minutagem_s"], "quem": lig["quem"], "fala": lig["fala"],
                    "trecho": lig["trecho"], "reformulacao": lig["reformulacao"], "tipo": lig["tipo"],
                    "checagem": c, "nota": {}, "motivo": {}})
                if c["nota"] > linha["nota"].get(base, -1):  # duas afirmações da mesma fala na mesma checagem
                    linha["nota"][base], linha["motivo"][base] = c["nota"], c["motivo"]
    return dict(sorted(por_chave.items(), key=lambda kv: (kv[1]["minutagem_s"], -max(kv[1]["nota"].values()))))


def _mmss(s):
    return f"{int(s // 60):02d}:{int(s % 60):02d}"


def planilha(avaliacao):
    """Grava avaliacao/marcacao.xlsx: uma linha por ligação proposta, com a coluna "marcação" para o
    operador escrever certa ou errada. Não sobrescreve uma planilha que já existe."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    arquivo = Path(avaliacao) / "marcacao.xlsx"
    if arquivo.exists():
        print(f"{arquivo} já existe; não sobrescrevo a marcação")
        return arquivo
    wb = Workbook()
    folha = wb.active
    folha.title = "ligações"
    folha.append(COLUNAS)
    for k, lin in ligacoes(avaliacao).items():
        c = lin["checagem"]
        motivos = list(dict.fromkeys(lin["motivo"].values()))
        motivo = motivos[0] if len(motivos) == 1 else "\n".join(f"{b}: {m}" for b, m in lin["motivo"].items())
        folha.append([_mmss(lin["minutagem_s"]), lin["quem"], lin["fala"], lin["trecho"], lin["reformulacao"],
                      lin["tipo"], c["agencia"], c["data"], c["rotulo"], c["frase"], c["url"],
                      lin["nota"].get("congelada"), lin["nota"].get("completa"), motivo, None, None, k])
        celula = folha.cell(folha.max_row, COLUNAS.index("link") + 1)
        celula.hyperlink, celula.style = c["url"], "Hyperlink"
    larguras = {"fala": 60, "trecho": 40, "reformulação": 40, "frase checada": 40, "link": 30, "motivo": 40,
                "observação": 30, "chave": 12}
    for i, nome in enumerate(COLUNAS, 1):
        letra = folha.cell(1, i).column_letter
        folha.column_dimensions[letra].width = larguras.get(nome, 12)
        folha.cell(1, i).font = Font(bold=True)
        for linha in range(2, folha.max_row + 1):
            folha.cell(linha, i).alignment = Alignment(wrap_text=True, vertical="top")
    marca = folha.cell(1, COLUNAS.index("marcação") + 1).column_letter
    dv = DataValidation(type="list", formula1='"certa,errada"', allow_blank=True)
    dv.add(f"{marca}2:{marca}{max(folha.max_row, 2)}")
    folha.add_data_validation(dv)
    folha.freeze_panes = "C2"
    folha.auto_filter.ref = folha.dimensions
    ajuda = wb.create_sheet("como marcar")
    for linha in [
        "Cada linha é uma ligação proposta pela sessão: uma fala do debate da Band e uma checagem candidata.",
        "Na coluna 'marcação', escreva certa se a checagem trata da mesma afirmação (mesmo sujeito, mesmo fato, "
        "mesmo número e período), ou errada se não trata.",
        "Não julgue se a afirmação é verdadeira: só se a ligação está certa.",
        "'nota congelada' e 'nota completa' são as notas dadas pelo Claude em cada rodada (vazia: a ligação não "
        "apareceu naquela rodada).",
        "Não mude a coluna 'chave'. Depois de marcar, salve e rode: "
        "python -P scripts/avaliacao_band.py relatorio",
    ]:
        ajuda.append([linha])
    ajuda.column_dimensions["A"].width = 120
    wb.save(arquivo)
    print(f"planilha: {folha.max_row - 1} ligações -> {arquivo}")
    return arquivo


# ---------- relatório ----------

MINIMO = 15  # ligações acima do limiar, todas certas, para liberar a publicação automática
FAIXAS = [(0.0, 0.5), (0.5, 0.6), (0.6, 0.7), (0.7, 0.8), (0.8, 0.9), (0.9, 1.0)]
DESLIGADA = "publicação automática desligada"


def le_marcas(avaliacao):
    """{chave: "certa" | "errada"} da planilha marcada, e as marcações que não se entendem."""
    import openpyxl

    folha = openpyxl.load_workbook(Path(avaliacao) / "marcacao.xlsx", read_only=True).worksheets[0]
    linhas = folha.iter_rows(values_only=True)
    cab = list(next(linhas))
    marcas, invalidas = {}, []
    for linha in linhas:
        chave, valor = linha[cab.index("chave")], linha[cab.index("marcação")]
        texto = str(valor or "").strip().lower()
        if not chave or not texto:
            continue
        if texto in ("certa", "certo", "c"):
            marcas[chave] = "certa"
        elif texto in ("errada", "errado", "e"):
            marcas[chave] = "errada"
        else:
            invalidas.append((chave, valor))
    return marcas, invalidas


def _conta(linhas):
    certas = sum(m == "certa" for _, m in linhas)
    erradas = sum(m == "errada" for _, m in linhas)
    return {"ligacoes": len(linhas), "certas": certas, "erradas": erradas,
            "sem_marcacao": len(linhas) - certas - erradas,
            "precisao": certas / (certas + erradas) if certas + erradas else None}


def escolhe_limiar(notas_e_marcas, minimo=MINIMO):
    """[(nota, marca)] → (menor nota t tal que as ligações com nota ≥ t estão todas marcadas certas e
    são pelo menos `minimo`, ou None; tabela por nota candidata)."""
    tabela, escolhido = [], None
    for t in sorted({n for n, _ in notas_e_marcas}):
        acima = _conta([(n, m) for n, m in notas_e_marcas if n >= t])
        tabela.append({"limiar": t, **acima})
        if escolhido is None and acima["certas"] == acima["ligacoes"] >= minimo:
            escolhido = t
    return escolhido, tabela


def relatorio(avaliacao, minimo=MINIMO, relogio=None):
    """Lê as rodadas e a planilha marcada; grava relatorio.md e limiar.json e devolve os números.
    Uma ligação conta com a maior nota que recebeu nas duas rodadas (é a que a publicaria sozinha)."""
    avaliacao = Path(avaliacao)
    agora = (relogio or (lambda: datetime.now(BRASILIA)))().isoformat(timespec="seconds")
    linhas = ligacoes(avaliacao)
    marcas, invalidas = le_marcas(avaliacao)
    por_base = {}
    for base in BASES:
        da_base = [(lin["nota"][base], marcas.get(k)) for k, lin in linhas.items() if base in lin["nota"]]
        faixas = [{"de": de, "ate": ate, **_conta([(n, m) for n, m in da_base
                                                   if de <= n < ate or (ate == 1.0 and n == 1.0)])}
                  for de, ate in FAIXAS]
        por_base[base] = {**_conta(da_base), "faixas": faixas}
    limiar, tabela = escolhe_limiar([(max(lin["nota"].values()), marcas.get(k)) for k, lin in linhas.items()],
                                    minimo)
    if limiar is None:
        origem = (f"avaliação da Band ({avaliacao / 'relatorio.md'}): nenhuma nota com zero erradas e pelo menos "
                  f"{minimo} ligações; {DESLIGADA}")
    else:
        n = next(t["ligacoes"] for t in tabela if t["limiar"] == limiar)
        origem = (f"avaliação da Band ({avaliacao / 'relatorio.md'}): {n} ligações com nota ≥ {limiar}, "
                  f"nenhuma errada, nas bases congelada e completa")
    (avaliacao / "limiar.json").write_text(json.dumps({"limiar": limiar, "origem": origem, "em": agora},
                                                      ensure_ascii=False, indent=1), encoding="utf-8")
    resultado = {"limiar": limiar, "origem": origem, "minimo": minimo, "ligacoes": len(linhas),
                 "sem_marcacao": sum(k not in marcas for k in linhas), "marcacoes_invalidas": invalidas,
                 "por_base": por_base, "tabela_limiares": tabela}
    (avaliacao / "relatorio.md").write_text(_relatorio_md(avaliacao, resultado, agora), encoding="utf-8")
    print((avaliacao / "relatorio.md").read_text(encoding="utf-8"))
    return resultado


def _pct(x):
    return "—" if x is None else f"{100 * x:.0f}%"


def _relatorio_md(avaliacao, r, agora):
    congelamento = avaliacao / "base_congelada" / "congelamento.json"
    c = json.loads(congelamento.read_text(encoding="utf-8")) if congelamento.exists() else None
    out = [f"# Avaliação da Band: relatório ({agora})", ""]
    if r["limiar"] is None:
        out += [f"**Decisão: {DESLIGADA}.** Nenhuma nota tem zero erradas e pelo menos {r['minimo']} ligações "
                "acima dela (ligações sem marcação não contam como certas).", ""]
    else:
        out += [f"**Decisão: limiar {r['limiar']}.** {r['origem']}.", "",
                "O limiar vale por checagem: só pode ir sozinha ao painel a checagem cuja própria nota é ≥ "
                "limiar (as demais checagens da mesma afirmação vão para a fila).", ""]
    out += [f"Gravado em `{avaliacao / 'limiar.json'}`.", ""]
    if c:
        out += ["## Bases", "",
                f"- completa: {c['checagens_na_base_completa']} checagens",
                f"- congelada (revisadas antes de {c['corte']}): {c['checagens_na_base_congelada']} checagens",
                f"  - excluídas sem data de revisão: {c['excluidas_sem_data_de_revisao']} (" +
                ", ".join(f"{ag} {n}" for ag, n in c.get("excluidas_sem_data_de_revisao_por_agencia", {}).items()) + ")",
                f"  - excluídas revisadas depois do corte (ou no dia do corte, que não tem hora): "
                f"{c['excluidas_revisadas_depois_do_corte']}",
                f"  - regra: {c['regra']}", ""]
    out += ["## Rodadas", "", "| base | afirmações | ligações propostas | sem par | descartadas | erros | chamadas | tokens (entrada/saída) |",
            "|---|---|---|---|---|---|---|---|"]
    for base in BASES:
        f = avaliacao / f"rodada_{base}" / "fila.json"
        s = json.loads(f.read_text(encoding="utf-8"))["resumo"]
        out.append(f"| {base} | {s['afirmacoes']} | {s['ligacoes_propostas']} | {s['sem_par']} | {s['descartados']} "
                   f"| {s['erros']} | {s['chamadas_claude']} | {s['tokens']['entrada']}/{s['tokens']['saida']} |")
    out += ["", "## Marcação", "",
            f"{r['ligacoes']} ligações distintas (fala × checagem) na planilha; {r['sem_marcacao']} sem marcação."]
    if r["marcacoes_invalidas"]:
        out.append("Marcações não reconhecidas (use certa ou errada): " +
                   "; ".join(f"{k}: {v!r}" for k, v in r["marcacoes_invalidas"]))
    for base in BASES:
        b = r["por_base"][base]
        out += ["", f"## Base {base}: {b['ligacoes']} ligações, precisão {_pct(b['precisao'])}", "",
                "| nota | ligações | certas | erradas | sem marcação | precisão |", "|---|---|---|---|---|---|"]
        out += [f"| {f['de']:.1f}–{f['ate']:.1f} | {f['ligacoes']} | {f['certas']} | {f['erradas']} | "
                f"{f['sem_marcacao']} | {_pct(f['precisao'])} |" for f in b["faixas"]]
    out += ["", "## Limiares candidatos (as duas bases, maior nota de cada ligação)", "",
            "| nota ≥ | ligações | certas | erradas | sem marcação |", "|---|---|---|---|---|"]
    out += [f"| {t['limiar']} | {t['ligacoes']} | {t['certas']} | {t['erradas']} | {t['sem_marcacao']} |"
            for t in r["tabela_limiares"]]
    return "\n".join(out) + "\n"


# ---------- tudo junto ----------

def avalia(dados, falas_json, saida=None, chamar=None, paralelo=3, origens=("revisao",), **kw):
    """Congela a base de DADOS/checagens, roda a sessão contra as duas versões e grava a planilha."""
    dados = Path(dados)
    saida = Path(saida) if saida else dados / "avaliacao"
    completa = dados / "checagens"
    if not (saida / "base_congelada" / "congelamento.json").exists():  # uma vez só: a rodada usa esta
        print(json.dumps(congela(completa, saida / "base_congelada", origens=origens), ensure_ascii=False, indent=1))
    falas = le_falas(falas_json)
    for base, pasta in (("congelada", saida / "base_congelada"), ("completa", completa)):
        rodada(falas, pasta, saida / f"rodada_{base}", chamar=chamar, paralelo=paralelo, nome_base=base, **kw)
    return planilha(saida)


class Duble:
    """Claude de mentira para ensaiar o comando sem gastar chamadas: em cada fala, a primeira frase vira
    a afirmação, e a primeira checagem da base que chegar entre as candidatas é confirmada com nota 0,5.
    Reconhece a fala e as candidatas pelo texto e pelas URLs que chegam no pedido, não pela forma dele.
    falas: textos das falas; urls: URLs das checagens da base."""

    def __init__(self, falas, urls):
        self.falas, self.urls = list(falas), list(urls)

    def __call__(self, etapa, prompt, conteudo, schema):
        if "afirmacoes" in schema["properties"]:
            fala = max((f for f in self.falas if f in conteudo), key=len, default="")
            frase = fala.split(". ")[0].strip()
            saida = {"afirmacoes": [{"trecho": frase, "reformulacao": frase, "tipo": "perene"}] if frase else []}
        else:
            url = next((u for u in self.urls if u in conteudo), None)
            saida = {"confirmadas": [{"url": url, "nota": 0.5, "motivo": "dublê"}] if url else []}
        return saida, {"etapa": etapa, "modelo_respondeu": "dublê", "prompt_sha256": "0" * 64,
                       "tokens": {"entrada": 0, "saida": 0}}


def main():
    import argparse

    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("comando", choices=["rodar", "relatorio"])
    ap.add_argument("--dados", help="pasta de dados (padrão: CHECAGEM_DADOS ou data/ do checkout principal)")
    ap.add_argument("--falas", help="falas de referência (padrão: DADOS/bloco1_final.json)")
    ap.add_argument("--saida", help="pasta da avaliação (padrão: DADOS/avaliacao)")
    ap.add_argument("--duble", action="store_true", help="ensaio com um Claude de mentira (sem chamadas)")
    ap.add_argument("--aceitar-publicacao", action="store_true",
                    help="na base congelada, aceitar também a data de publicação (UOL/BOL não têm data de revisão)")
    ap.add_argument("--paralelo", type=int, default=3, help="falas analisadas ao mesmo tempo")
    ap.add_argument("--minimo", type=int, default=MINIMO, help="ligações acima do limiar, todas certas")
    a = ap.parse_args()
    dados = pasta_dados(a.dados)
    saida = Path(a.saida) if a.saida else dados / "avaliacao"
    if a.comando == "rodar":
        if not a.duble and not os.environ.get("ANTHROPIC_API_KEY"):  # num worktree o .env só existe no checkout principal
            env = raiz_principal() / ".env"
            for linha in env.read_text(encoding="utf-8").splitlines() if env.exists() else []:
                nome, _, valor = linha.partition("=")
                if nome.strip() == "ANTHROPIC_API_KEY" and valor.strip():
                    os.environ["ANTHROPIC_API_KEY"] = valor.strip()
        falas = a.falas or dados / "bloco1_final.json"
        duble = None
        if a.duble:
            from base_checagens import le_base

            duble = Duble([x for _, _, x in le_falas(falas)], [r["url"] for r in le_base(dados / "checagens")])
        avalia(dados, falas, saida, chamar=duble,
               paralelo=a.paralelo, origens=("revisao", "publicacao") if a.aceitar_publicacao else ("revisao",))
    else:
        relatorio(saida, minimo=a.minimo)


if __name__ == "__main__":
    main()
