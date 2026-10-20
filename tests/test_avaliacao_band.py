"""Avaliação da Band: base congelada, duas rodadas da sessão, planilha de marcação e relatório.

O Claude é trocado por um dublê; a busca por sentido roda de verdade sobre a base de exemplo de
tests/fixtures/base_sessao.json mais checagens publicadas depois do debate da Band.
"""
import json
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from avaliacao_band import avalia, congela, relatorio
from base_checagens import Indice, atualiza_indice

FIXTURES = Path(__file__).parent / "fixtures"
BRASILIA = timezone(timedelta(hours=-3))

DESEMPREGO_2025 = "https://example.org/aosfatos/desemprego-2025"
DESEMPREGO_BAND = "https://example.org/lupa/debate-band-desemprego"  # revisada depois do debate
PIX_ESTADAO = "https://example.org/estadao/pix-bolsonaro"
PIX_AFP = "https://example.org/afp/pix-criacao"

FALA_DESEMPREGO = ("Olha, o meu governo gerou emprego como nunca. Hoje o desemprego é o menor da história "
                   "do IBGE, e isso ninguém pode negar.")
FALA_PIX = "Quem criou o Pix foi o governo do meu pai, isso todo mundo sabe."
FALA_MEDIACAO = "Candidato Lula, o senhor tem um minuto para a resposta sobre o desemprego."

AFIRMACOES = {
    FALA_DESEMPREGO: {"trecho": "o desemprego é o menor da história do IBGE",
                      "reformulacao": "A taxa de desemprego é a menor da série histórica do IBGE.", "tipo": "numero",
                      "confirma": {DESEMPREGO_2025: (0.92, "Mesmo indicador e mesma série."),
                                   DESEMPREGO_BAND: (0.97, "A mesma frase dita no debate da Band.")}},
    FALA_PIX: {"trecho": "Quem criou o Pix foi o governo do meu pai",
               "reformulacao": "O Pix foi criado no governo de Jair Bolsonaro.", "tipo": "perene",
               "confirma": {PIX_ESTADAO: (0.95, "Mesma afirmação sobre a criação do Pix."),
                            PIX_AFP: (0.7, "Mesma afirmação, atribuída a redes sociais.")}},
}


class Claude:
    """Dublê do Claude com respostas prontas no formato da saída estruturada: confirma só as
    checagens de AFIRMACOES que chegarem entre as candidatas."""

    def __init__(self):
        self.chamadas = 0

    def __call__(self, etapa, prompt, conteudo, schema):
        self.chamadas += 1
        if prompt == "extrair_afirmacoes.md":
            achadas = [a for fala, a in AFIRMACOES.items() if fala in conteudo]
            saida = {"afirmacoes": [{k: a[k] for k in ("trecho", "reformulacao", "tipo")} for a in achadas]}
        else:
            [a] = [a for a in AFIRMACOES.values() if a["trecho"] in conteudo]
            saida = {"confirmadas": [{"url": u, "nota": n, "motivo": m}
                                     for u, (n, m) in a["confirma"].items() if f"URL: {u}\n" in conteudo]}
        return saida, {"etapa": etapa, "modelo_respondeu": "dublê", "prompt_sha256": "0" * 64,
                       "tokens": {"entrada": 0, "saida": 0}}


@pytest.fixture(scope="module")
def dados(tmp_path_factory):
    """Pasta de dados com a base completa (a de exemplo mais uma checagem revisada na noite seguinte
    ao debate da Band) e as falas de referência."""
    dados = tmp_path_factory.mktemp("dados")
    registros = json.loads((FIXTURES / "base_sessao.json").read_text(encoding="utf-8"))["registros"]
    registros.append({**registros[0], "id": DESEMPREGO_BAND, "url": DESEMPREGO_BAND, "agencia": "Lupa",
                      "frase": "Lula disse no debate da Band que o desemprego é o menor da história do IBGE",
                      "data": "2026-10-19", "data_origem": "revisao", "rotulo": "Verdadeiro"})
    pasta = dados / "checagens"
    pasta.mkdir()
    (pasta / "base.json").write_text(json.dumps({"registros": registros}, ensure_ascii=False), encoding="utf-8")
    atualiza_indice(pasta, registros)
    turnos = [("Mediação", 16226.0, FALA_MEDIACAO), ("Lula", 16230.5, FALA_DESEMPREGO), ("Flávio", 16290.0, FALA_PIX)]
    (dados / "falas.json").write_text(json.dumps({"turns": [{"speaker": q, "start": t, "text": x}
                                                            for q, t, x in turnos]}, ensure_ascii=False),
                                      encoding="utf-8")
    return dados


@pytest.fixture(scope="module")
def avaliacao(dados):
    avalia(dados, dados / "falas.json", chamar=Claude())
    return dados / "avaliacao"


def linhas(xlsx):
    folha = openpyxl.load_workbook(xlsx).active
    cab, *resto = folha.iter_rows(values_only=True)
    return [dict(zip(cab, linha)) for linha in resto]


def registro(url, data, origem, frase="Uma frase checada qualquer"):
    return {"id": url, "url": url, "agencia": "Aos Fatos", "frase": frase, "quem_disse": None, "data": data,
            "data_origem": origem, "rotulo": "Falso"}


def grava_base(pasta, registros, vetores):
    pasta.mkdir(parents=True)
    (pasta / "base.json").write_text(json.dumps({"registros": registros}), encoding="utf-8")
    np.save(pasta / "vetores.npy", np.asarray(vetores, dtype=np.float32))
    (pasta / "vetores.json").write_text(json.dumps({"ids": [r["id"] for r in registros],
                                                    "frases": [r["frase"] for r in registros]}), encoding="utf-8")


def test_base_congelada_so_tem_checagens_revisadas_antes_do_debate_da_band(tmp_path):
    registros = [
        registro("https://a/antes", "2026-10-10", "revisao"),
        registro("https://a/dia-do-debate", "2026-10-18", "revisao"),
        registro("https://a/noite-do-debate", "2026-10-19", "revisao"),  # dia UTC que começa 21h de 18/10
        registro("https://a/depois", "2026-10-20", "revisao"),
        registro("https://a/so-publicacao", "2025-03-01", "publicacao"),
        registro("https://a/so-afirmacao", "2024-01-01", "afirmacao"),
        registro("https://a/sem-data", None, None),
    ]
    vetores = np.eye(len(registros), 8)
    grava_base(tmp_path / "completa", registros, vetores)

    resumo = congela(tmp_path / "completa", tmp_path / "congelada")

    congelada = Indice.da_pasta(tmp_path / "congelada")
    assert [r["id"] for r in congelada.registros] == ["https://a/antes", "https://a/dia-do-debate"]
    assert np.array_equal(congelada.X, vetores[:2])
    assert resumo["checagens_na_base_completa"] == 7
    assert resumo["checagens_na_base_congelada"] == 2
    assert resumo["excluidas_sem_data_de_revisao"] == 3
    assert resumo["excluidas_revisadas_depois_do_corte"] == 2
    assert json.loads((tmp_path / "congelada" / "congelamento.json").read_text(encoding="utf-8")) == resumo
    assert resumo["excluidas_sem_data_de_revisao_por_agencia"] == {"Aos Fatos": 3}


def test_base_congelada_pode_aceitar_tambem_a_data_de_publicacao(tmp_path):
    # UOL e BOL nunca trazem data de revisão; a data de publicação vem da URL
    registros = [registro("https://a/revisada", "2026-10-10", "revisao"),
                 registro("https://uol/2025/03/01/x", "2025-03-01", "publicacao"),
                 registro("https://uol/2026/10/20/y", "2026-10-20", "publicacao"),
                 registro("https://estadao/z", "2024-01-01", "afirmacao")]
    grava_base(tmp_path / "completa", registros, np.eye(4, 8))

    resumo = congela(tmp_path / "completa", tmp_path / "congelada", origens=("revisao", "publicacao"))

    assert [r["id"] for r in Indice.da_pasta(tmp_path / "congelada").registros] == [
        "https://a/revisada", "https://uol/2025/03/01/x"]
    assert (resumo["excluidas_sem_data_de_revisao"], resumo["excluidas_revisadas_depois_do_corte"]) == (1, 1)


def test_duas_rodadas_guardadas_em_separado_com_o_registro_da_sessao(avaliacao):
    congelada = json.loads((avaliacao / "rodada_congelada" / "fila.json").read_text(encoding="utf-8"))
    completa = json.loads((avaliacao / "rodada_completa" / "fila.json").read_text(encoding="utf-8"))

    assert completa["resumo"]["checagens_na_base"] == 21
    assert congelada["resumo"]["checagens_na_base"] < 21
    assert [lig["quem"] for lig in completa["fila"]] == ["Lula", "Flávio"]
    for rodada in ("rodada_congelada", "rodada_completa"):
        eventos = (avaliacao / rodada / "eventos.jsonl").read_text(encoding="utf-8").splitlines()
        chamadas = [json.loads(e)["chamada"] for e in eventos if '"chamada"' in e]
        assert chamadas and all(c["modelo_respondeu"] == "dublê" for c in chamadas if c)


def test_planilha_tem_uma_linha_por_ligacao_e_as_iguais_nas_duas_rodadas_aparecem_uma_vez(avaliacao):
    por_link = {lin["link"]: lin for lin in linhas(avaliacao / "marcacao.xlsx")}

    assert set(por_link) == {DESEMPREGO_2025, DESEMPREGO_BAND, PIX_ESTADAO, PIX_AFP}
    desemprego = por_link[DESEMPREGO_2025]
    assert (desemprego["quem"], desemprego["fala"]) == ("Lula", FALA_DESEMPREGO)
    assert desemprego["trecho"] == "o desemprego é o menor da história do IBGE"
    assert desemprego["reformulação"] == "A taxa de desemprego é a menor da série histórica do IBGE."
    assert desemprego["agência"] == "Aos Fatos"
    assert desemprego["frase checada"] == "Taxa de desemprego no Brasil é a menor da série histórica do IBGE"
    assert (desemprego["nota congelada"], desemprego["nota completa"]) == (0.92, 0.92)
    assert desemprego["motivo"] == "Mesmo indicador e mesma série."
    assert desemprego["marcação"] is None
    # a checagem revisada depois do debate só existe na base completa
    assert (por_link[DESEMPREGO_BAND]["nota congelada"], por_link[DESEMPREGO_BAND]["nota completa"]) == (None, 0.97)


# ---------- relatório, depois da marcação do operador ----------

def marcada(avaliacao, tmp_path, marcas):
    """Cópia da avaliação com a planilha marcada como o operador faria no Excel: marcas = {link: certa|errada}."""
    pasta = tmp_path / "avaliacao"
    shutil.copytree(avaliacao, pasta)
    wb = openpyxl.load_workbook(pasta / "marcacao.xlsx")
    folha = wb.active
    cab = [c.value for c in folha[1]]
    for linha in folha.iter_rows(min_row=2):
        link = linha[cab.index("link")].value
        if link in marcas:
            linha[cab.index("marcação")].value = marcas[link]
    wb.save(pasta / "marcacao.xlsx")
    return pasta


def limiar_gravado(pasta):
    return json.loads((pasta / "limiar.json").read_text(encoding="utf-8"))


TODAS_CERTAS = {DESEMPREGO_2025: "certa", DESEMPREGO_BAND: "certa", PIX_ESTADAO: "certa", PIX_AFP: "certa"}


def test_com_todas_certas_o_limiar_e_a_menor_nota(avaliacao, tmp_path):
    pasta = marcada(avaliacao, tmp_path, TODAS_CERTAS)

    relatorio(pasta, minimo=4, relogio=lambda: datetime(2026, 10, 21, 10, 0, tzinfo=BRASILIA))

    gravado = limiar_gravado(pasta)
    assert gravado["limiar"] == 0.7
    assert gravado["em"] == "2026-10-21T10:00:00-03:00"
    assert "4 ligações" in gravado["origem"]


def test_o_limiar_fica_acima_da_nota_de_qualquer_ligacao_errada(avaliacao, tmp_path):
    pasta = marcada(avaliacao, tmp_path, {**TODAS_CERTAS, DESEMPREGO_2025: "errada", PIX_AFP: "errada"})

    relatorio(pasta, minimo=2)

    assert limiar_gravado(pasta)["limiar"] == 0.95


def test_ligacao_sem_marcacao_nao_conta_como_certa(avaliacao, tmp_path):
    pasta = marcada(avaliacao, tmp_path, {**TODAS_CERTAS, PIX_AFP: None})

    r = relatorio(pasta, minimo=2)

    assert limiar_gravado(pasta)["limiar"] == 0.92
    assert r["sem_marcacao"] == 1


def test_sem_ligacoes_suficientes_a_publicacao_automatica_fica_desligada(avaliacao, tmp_path):
    pasta = marcada(avaliacao, tmp_path, TODAS_CERTAS)

    relatorio(pasta)  # mínimo padrão: ~15 ligações

    gravado = limiar_gravado(pasta)
    assert gravado["limiar"] is None
    assert "publicação automática desligada" in gravado["origem"]
    assert "publicação automática desligada" in (pasta / "relatorio.md").read_text(encoding="utf-8")


def test_relatorio_traz_ligacoes_e_precisao_por_faixa_de_nota_em_cada_base(avaliacao, tmp_path):
    pasta = marcada(avaliacao, tmp_path, {**TODAS_CERTAS, PIX_ESTADAO: "errada"})

    r = relatorio(pasta)

    congelada, completa = r["por_base"]["congelada"], r["por_base"]["completa"]
    assert (congelada["ligacoes"], completa["ligacoes"]) == (3, 4)
    faixa_alta = next(f for f in completa["faixas"] if f["de"] == 0.9)
    assert (faixa_alta["ligacoes"], faixa_alta["certas"], faixa_alta["erradas"]) == (3, 2, 1)
    assert faixa_alta["precisao"] == pytest.approx(2 / 3)
    texto = (pasta / "relatorio.md").read_text(encoding="utf-8")
    assert "sem data de revisão" in texto


def test_rodada_em_que_todas_as_chamadas_falham_nao_fica_valendo(dados, tmp_path):
    def sem_chave(etapa, prompt, conteudo, schema):
        raise TypeError("Could not resolve authentication method")

    with pytest.raises(SystemExit, match="authentication"):
        avalia(dados, dados / "falas.json", saida=tmp_path / "avaliacao", chamar=sem_chave)

    assert not (tmp_path / "avaliacao" / "rodada_congelada" / "fila.json").exists()
    assert not (tmp_path / "avaliacao" / "marcacao.xlsx").exists()


def test_rodar_de_novo_nao_chama_o_claude_nem_apaga_a_marcacao(dados, avaliacao, tmp_path):
    pasta = marcada(avaliacao, tmp_path, TODAS_CERTAS)
    claude = Claude()

    avalia(dados, dados / "falas.json", saida=pasta, chamar=claude)

    assert claude.chamadas == 0
    assert {lin["link"]: lin["marcação"] for lin in linhas(pasta / "marcacao.xlsx")} == TODAS_CERTAS
